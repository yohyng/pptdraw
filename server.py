#!/usr/bin/env python3
"""
PPTX Pen Annotator - Backend Merge Server
Receives transparent PNG annotations and merges them into a .pptx file.

Usage:
    PPTX_PATH=./presentation.pptx python server.py
    PORT=5000 PPTX_PATH=./my.pptx python server.py
"""

import base64
import io
import os
from pathlib import Path

from flask import Flask, request, jsonify
from flask_cors import CORS
from pptx import Presentation

app = Flask(__name__)
CORS(app)

PPTX_PATH = os.environ.get("PPTX_PATH", "presentation.pptx")


@app.route("/merge", methods=["POST"])
def merge():
    data = request.get_json()
    if not data or "annotations" not in data:
        return jsonify({"error": "No annotations provided"}), 400

    path = Path(PPTX_PATH)
    if not path.exists():
        return jsonify({"error": f"PPTX not found: {PPTX_PATH}"}), 404

    prs = Presentation(str(path))
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    processed = 0

    for annotation in data["annotations"]:
        slide_index = annotation.get("slide_index", 0)
        png_b64 = annotation.get("png_base64", "")

        if not png_b64:
            continue
        if slide_index < 0 or slide_index >= len(prs.slides):
            continue

        png_bytes = base64.b64decode(png_b64)
        slide = prs.slides[slide_index]
        slide.shapes.add_picture(
            io.BytesIO(png_bytes),
            left=0,
            top=0,
            width=slide_width,
            height=slide_height,
        )
        processed += 1

    out_path = path.with_stem(path.stem + "_annotated")
    prs.save(str(out_path))

    return jsonify({
        "message": f"{processed}枚のスライドに注釈を追加しました",
        "processed": processed,
        "output_file": str(out_path),
    })


@app.route("/status", methods=["GET"])
def status():
    path = Path(PPTX_PATH)
    slide_count = 0
    if path.exists():
        try:
            prs = Presentation(str(path))
            slide_count = len(prs.slides)
        except Exception:
            pass
    return jsonify({
        "ok": True,
        "pptx_path": str(path),
        "pptx_exists": path.exists(),
        "slide_count": slide_count,
    })


@app.route("/slides", methods=["GET"])
def list_slides():
    path = Path(PPTX_PATH)
    if not path.exists():
        return jsonify({"error": "PPTX not found"}), 404
    prs = Presentation(str(path))
    return jsonify({
        "slide_count": len(prs.slides),
        "slide_width_emu": prs.slide_width,
        "slide_height_emu": prs.slide_height,
    })


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    debug = os.environ.get("DEBUG", "false").lower() == "true"
    print(f"Starting PPTX Merge Server on port {port}")
    print(f"Target PPTX: {PPTX_PATH}")
    app.run(host="0.0.0.0", port=port, debug=debug)
